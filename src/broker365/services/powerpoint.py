"""PowerPoint (.pptx). Graph has no content API for PowerPoint, so every edit is a round-trip:
download -> edit with python-pptx -> upload with If-Match eTag."""

from __future__ import annotations

from collections.abc import Iterator
from typing import Any, Literal

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Inches

from broker365.errors import Broker365Error
from broker365.services import files, tool
from broker365.services.word import replace_in_paragraph

PPT_EXTS = (".pptx",)  # python-pptx rejects macro/template content types
# kind -> (layout name, index in the default template as a fallback)
LAYOUTS = {
    "title": ("Title Slide", 0),
    "title_content": ("Title and Content", 1),
    "blank": ("Blank", 6),
}


# ---- helpers ---------------------------------------------------------------------------


def _load(item: str) -> tuple[Any, dict[str, Any]]:
    data, meta = files.download_office(item, PPT_EXTS, "PowerPoint")
    return files.parse_office(data, Presentation, "PowerPoint"), meta


def _layout(prs: Any, kind: str) -> Any:
    if kind not in LAYOUTS:
        raise Broker365Error(f"Unknown layout '{kind}' (use title, title_content or blank)")
    name, index = LAYOUTS[kind]
    layouts = prs.slide_layouts
    for layout in layouts:
        if layout.name == name:
            return layout
    return layouts[min(index, len(layouts) - 1)]


def _body(slide: Any) -> Any | None:
    """First non-title placeholder that holds text (subtitle or content body)."""
    for ph in slide.placeholders:
        if ph.placeholder_format.idx != 0 and ph.has_text_frame:
            return ph
    return None


def _fill_bullets(slide: Any, prs: Any, bullets: list[str]) -> None:
    body = _body(slide)
    if body is None:
        top = Inches(1.5) if slide.shapes.title is not None else Inches(0.5)
        body = slide.shapes.add_textbox(
            Inches(0.5), top, prs.slide_width - Inches(1), prs.slide_height - top - Inches(0.5)
        )
    tf = body.text_frame
    for i, raw in enumerate(bullets):
        text = raw.rstrip()
        stripped = text.lstrip(" ")
        level = min((len(text) - len(stripped)) // 2, 8)
        if stripped[:2] in ("- ", "* ", "• "):
            stripped = stripped[2:]
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text, p.level = stripped, level


def _add_slide(
    prs: Any,
    title: str | None,
    bullets: list[str] | None = None,
    notes: str | None = None,
    layout: str = "title_content",
    subtitle: str | None = None,
) -> Any:
    slide = prs.slides.add_slide(_layout(prs, layout))
    if title:
        if slide.shapes.title is not None:
            slide.shapes.title.text = title
        else:
            box = slide.shapes.add_textbox(
                Inches(0.5), Inches(0.3), prs.slide_width - Inches(1), Inches(1)
            )
            box.text_frame.text = title
    if subtitle and (body := _body(slide)) is not None:
        body.text_frame.text = subtitle
    elif bullets:
        _fill_bullets(slide, prs, bullets)
    if notes:
        slide.notes_slide.notes_text_frame.text = notes
    return slide


def _shapes(shapes: Any) -> Iterator[Any]:
    for shape in shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from _shapes(shape.shapes)
        else:
            yield shape


def _paragraphs(shape: Any) -> Iterator[Any]:
    if shape.has_text_frame:
        yield from shape.text_frame.paragraphs
    if getattr(shape, "has_table", False):
        for row in shape.table.rows:
            for cell in row.cells:
                yield from cell.text_frame.paragraphs


def _shape_text(shape: Any) -> str:
    return "\n".join(p.text for p in _paragraphs(shape))


# ---- tools ----------------------------------------------------------------------------


@tool("files", writes=True)
def ppt_create(
    remote_path: str,
    slides: list[dict[str, Any]],
    conflict: Literal["fail", "replace", "rename"] = "fail",
) -> dict[str, Any]:
    """Create a .pptx at an absolute OneDrive path from an outline. Each slide is
    {"title", "subtitle"?, "bullets"? (two-space indent = sub-bullet), "notes"?,
    "layout"?: "title"|"title_content"|"blank"}."""
    if not slides:
        raise Broker365Error("Give at least one slide")
    prs = Presentation()
    for i, s in enumerate(slides):
        if not isinstance(s, dict):
            raise Broker365Error(f"Slide {i + 1} must be an object with a 'title'")
        default = "title" if i == 0 and s.get("subtitle") else "title_content"
        _add_slide(
            prs,
            s.get("title"),
            s.get("bullets"),
            s.get("notes"),
            s.get("layout") or default,
            s.get("subtitle"),
        )
    created = files.create_bytes(remote_path, files.to_bytes(prs), conflict)
    return {**created, "slides": len(slides)}


@tool("files")
def ppt_read(item: str) -> dict[str, Any]:
    """Read a .pptx: per slide its title, other text (incl. tables) and speaker notes.
    `item` is an absolute OneDrive path ('/Docs/deck.pptx') or item id."""
    prs, meta = _load(item)
    out = []
    for n, slide in enumerate(prs.slides, 1):
        title = slide.shapes.title
        texts = [
            _shape_text(s)
            for s in _shapes(slide.shapes)
            if title is None or s.shape_id != title.shape_id
        ]
        notes = slide.notes_slide.notes_text_frame.text if slide.has_notes_slide else ""
        out.append(
            {
                "index": n,
                "title": title.text if title is not None else "",
                "text": "\n".join(t for t in texts if t.strip()),
                "notes": notes,
            }
        )
    return {**files.summarize(meta), "slides": out}


@tool("files", writes=True)
def ppt_add_slide(
    item: str,
    title: str,
    bullets: list[str] | None = None,
    notes: str | None = None,
    position: int | None = None,
) -> dict[str, Any]:
    """Add a title+bullets slide to a .pptx (`item` = absolute OneDrive path or item id),
    at the end or at 1-based `position`. Download->edit->upload; existing complex content
    like SmartArt is preserved but not editable."""
    prs, meta = _load(item)
    _add_slide(prs, title, bullets, notes)
    count = len(prs.slides)
    index = count
    if position is not None:
        index = max(1, min(position, count))
        ids = prs.slides._sldIdLst
        el = ids[-1]
        ids.remove(el)
        ids.insert(index - 1, el)
    saved = files.replace_bytes(meta, files.to_bytes(prs))
    return {**saved, "slide_index": index, "slides": count}


@tool("files", writes=True)
def ppt_replace_text(item: str, find: str, replace: str) -> dict[str, Any]:
    """Find/replace text (case-sensitive) in every slide's shapes and table cells of a .pptx
    (`item` = absolute OneDrive path or item id). Download->edit->upload; SmartArt and charts
    are preserved but not searched. A match split across differently formatted runs takes
    the formatting of the paragraph's first run."""
    if not find:
        raise Broker365Error("'find' must not be empty")
    prs, meta = _load(item)
    count = sum(
        replace_in_paragraph(p, find, replace)
        for slide in prs.slides
        for shape in _shapes(slide.shapes)
        for p in _paragraphs(shape)
    )
    return files.save_if_changed(meta, count > 0, lambda: files.to_bytes(prs), replacements=count)


@tool("files")
def ppt_export_pdf(item: str, dest: str | None = None, overwrite: bool = False) -> dict[str, Any]:
    """Convert a .pptx (absolute OneDrive path or item id) to PDF and save it locally."""
    return files.files_convert(item, "pdf", dest, overwrite)
