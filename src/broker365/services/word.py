"""Word (.docx). Graph has no content API for Word, so every edit is a round-trip:
download -> edit with python-docx -> upload with If-Match eTag."""

from __future__ import annotations

import re
from collections.abc import Iterator
from typing import Any, Literal

import docx
from docx.document import Document
from docx.shared import Pt
from docx.table import Table

from broker365.errors import Broker365Error
from broker365.services import files, tool

_HEADING = re.compile(r"^(#{1,6})\s+(.*?)\s*#*\s*$")
_BULLET = re.compile(r"^(\s*)[-*+]\s+(.*)$")
_NUMBERED = re.compile(r"^(\s*)\d+[.)]\s+(.*)$")
_RULE = re.compile(r"^\s*([-*_])(\s*\1){2,}\s*$")
_TABLE_SEP = re.compile(r"^\s*\|?\s*:?-+:?\s*(\|\s*:?-+:?\s*)*\|?\s*$")
_INLINE = re.compile(r"(\*\*.+?\*\*|__.+?__|`[^`]+`|\*[^*\s][^*]*?\*)")
_LINK = re.compile(r"\[([^\]]+)\]\(([^)]+)\)")
_PLACEHOLDER = re.compile(r"\{\{\s*([^{}]+?)\s*\}\}")
MONO = "Consolas"
WORD_EXTS = (".docx",)  # python-docx rejects macro/template content types


# ---- markdown -> docx -----------------------------------------------------------------


def _style(doc: Document, name: str) -> str | None:
    return name if name in doc.styles else None


def _add_runs(paragraph: Any, text: str) -> None:
    for part in _INLINE.split(_LINK.sub(r"\1 (\2)", text)):
        if not part:
            continue
        if part[:2] in ("**", "__") and part[-2:] == part[:2] and len(part) > 4:
            paragraph.add_run(part[2:-2]).bold = True
        elif part.startswith("`") and part.endswith("`") and len(part) > 2:
            paragraph.add_run(part[1:-1]).font.name = MONO
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            paragraph.add_run(part[1:-1]).italic = True
        else:
            paragraph.add_run(part)


def _add_list_item(doc: Document, base: str, indent: str, text: str) -> None:
    level = min(len(indent.expandtabs(4)) // 2, 2)
    name = _style(doc, f"{base} {level + 1}" if level else base) or _style(doc, base)
    p = doc.add_paragraph(style=name)
    if name is None:
        p.add_run("• " if base == "List Bullet" else "- ")
    _add_runs(p, text)


def _split_row(line: str) -> list[str]:
    return [c.strip() for c in line.strip().strip("|").split("|")]


def _add_table(doc: Document, lines: list[str]) -> None:
    rows = [_split_row(ln) for ln in lines if not _TABLE_SEP.match(ln)]
    has_header = len(lines) > 1 and bool(_TABLE_SEP.match(lines[1]))
    width = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=width)
    if style := _style(doc, "Table Grid"):
        table.style = style
    for r, row in enumerate(rows):
        for c, text in enumerate(row):
            p = table.cell(r, c).paragraphs[0]
            _add_runs(p, text)
            if has_header and r == 0:
                for run in p.runs:
                    run.bold = True


def _add_markdown(doc: Document, markdown: str) -> None:
    """Append a practical Markdown subset: headings, paragraphs, bullet/numbered lists,
    **bold**, *italic*, `code`, pipe tables and fenced code blocks."""
    lines = markdown.replace("\r\n", "\n").split("\n")
    para: list[str] = []

    def flush() -> None:
        if para:
            _add_runs(doc.add_paragraph(), " ".join(s.strip() for s in para))
            para.clear()

    i = 0
    while i < len(lines):
        line = lines[i]
        if line.strip().startswith("```"):
            flush()
            i += 1
            while i < len(lines) and not lines[i].strip().startswith("```"):
                p = doc.add_paragraph(style=_style(doc, "No Spacing"))
                run = p.add_run(lines[i])
                run.font.name, run.font.size = MONO, Pt(9.5)
                i += 1
        elif line.lstrip().startswith("|"):
            flush()
            block = []
            while i < len(lines) and lines[i].lstrip().startswith("|"):
                block.append(lines[i])
                i += 1
            _add_table(doc, block)
            continue
        elif not line.strip() or _RULE.match(line):
            flush()
        elif m := _HEADING.match(line):
            flush()
            level = len(m.group(1))
            p = doc.add_paragraph(style=_style(doc, f"Heading {level}"))
            _add_runs(p, m.group(2))
            if p.style.name != f"Heading {level}":
                for run in p.runs:
                    run.bold = True
        elif m := _BULLET.match(line):
            flush()
            _add_list_item(doc, "List Bullet", m.group(1), m.group(2))
        elif m := _NUMBERED.match(line):
            flush()
            _add_list_item(doc, "List Number", m.group(1), m.group(2))
        else:
            para.append(line)
        i += 1
    flush()


# ---- text replacement (shared with powerpoint) -----------------------------------------


def replace_in_paragraph(paragraph: Any, find: str, replace: str, match_case: bool = True) -> int:
    """Replace `find` in a python-docx or python-pptx paragraph; returns the count.
    Matches inside a single run keep their formatting. A match split across runs collapses
    the paragraph's text into its first run (that run's formatting wins)."""
    runs = list(paragraph.runs)
    if not find or not runs:
        return 0
    pattern = re.compile(re.escape(find), 0 if match_case else re.IGNORECASE)
    total = len(pattern.findall("".join(r.text for r in runs)))
    if not total:
        return 0
    if sum(len(pattern.findall(r.text)) for r in runs) == total:
        for r in runs:
            if pattern.search(r.text):
                r.text = pattern.sub(lambda _: replace, r.text)
    else:
        runs[0].text = pattern.sub(lambda _: replace, "".join(r.text for r in runs))
        for r in runs[1:]:
            r.text = ""
    return total


def _table_paragraphs(table: Table) -> Iterator[Any]:
    seen: set[int] = set()
    for row in table.rows:
        for cell in row.cells:
            if id(cell._tc) in seen:  # merged cells repeat
                continue
            seen.add(id(cell._tc))
            yield from cell.paragraphs
            for inner in cell.tables:
                yield from _table_paragraphs(inner)


def _all_paragraphs(doc: Document) -> Iterator[Any]:
    """Body, table cells, and every defined header/footer."""
    yield from doc.paragraphs
    for t in doc.tables:
        yield from _table_paragraphs(t)
    seen: set[int] = set()
    for s in doc.sections:
        for part in (
            s.header,
            s.footer,
            s.first_page_header,
            s.first_page_footer,
            s.even_page_header,
            s.even_page_footer,
        ):
            if part.is_linked_to_previous or id(part.part) in seen:
                continue
            seen.add(id(part.part))
            yield from part.paragraphs
            for t in part.tables:
                yield from _table_paragraphs(t)


def _replace_all(doc: Document, find: str, replace: str, match_case: bool = True) -> int:
    return sum(replace_in_paragraph(p, find, replace, match_case) for p in _all_paragraphs(doc))


# ---- helpers ---------------------------------------------------------------------------


def _load(item: str) -> tuple[Document, dict[str, Any]]:
    data, meta = files.download_office(item, WORD_EXTS, "Word")
    return files.parse_office(data, docx.Document, "Word"), meta


def _blocks(doc: Document) -> Iterator[tuple[str, Any, int]]:
    """(kind, value, size) for word_read: non-empty paragraphs, then tables."""
    for p in doc.paragraphs:
        if p.text.strip():
            style = p.style.name if p.style is not None else ""
            yield "paragraphs", {"style": style, "text": p.text}, len(p.text)
    for t in doc.tables:
        rows = [[c.text for c in r.cells] for r in t.rows]
        yield "tables", rows, sum(len(c) for r in rows for c in r)


# ---- tools ----------------------------------------------------------------------------


@tool("files", writes=True)
def word_create(
    remote_path: str, markdown: str, conflict: Literal["fail", "replace", "rename"] = "fail"
) -> dict[str, Any]:
    """Create a .docx at an absolute OneDrive path from Markdown (headings, paragraphs,
    -/1. lists, **bold**, *italic*, `code`, pipe tables, ``` code blocks)."""
    doc = docx.Document()
    _add_markdown(doc, markdown)
    return files.create_bytes(remote_path, files.to_bytes(doc), conflict)


@tool("files")
def word_read(item: str, max_chars: int = 100_000) -> dict[str, Any]:
    """Read a .docx as paragraphs (with style names) and tables. `item` is an absolute
    OneDrive path ('/Docs/report.docx') or item id."""
    doc, meta = _load(item)
    out: dict[str, list[Any]] = {"paragraphs": [], "tables": []}
    room, truncated = max_chars, False
    for kind, value, size in _blocks(doc):
        if size > room:
            truncated = True
            break
        room -= size
        out[kind].append(value)
    return {**files.summarize(meta), **out, "truncated": truncated}


@tool("files", writes=True)
def word_append(item: str, markdown: str) -> dict[str, Any]:
    """Append Markdown to the end of an existing .docx (`item` = absolute OneDrive path or
    item id). Download->edit->upload; SmartArt and other complex content is kept as-is."""
    doc, meta = _load(item)
    _add_markdown(doc, markdown)
    return files.replace_bytes(meta, files.to_bytes(doc))


@tool("files", writes=True)
def word_replace_text(
    item: str, find: str, replace: str, match_case: bool = True
) -> dict[str, Any]:
    """Find/replace text in a .docx body, tables, headers and footers (`item` = absolute
    OneDrive path or item id). Download->edit->upload; complex content like SmartArt or
    text boxes is preserved but not searched. A match split across differently formatted
    runs takes the formatting of the paragraph's first run."""
    if not find:
        raise Broker365Error("'find' must not be empty")
    doc, meta = _load(item)
    count = _replace_all(doc, find, replace, match_case)
    return files.save_if_changed(meta, count > 0, lambda: files.to_bytes(doc), replacements=count)


@tool("files", writes=True)
def word_fill_template(
    item: str, values: dict[str, str], save_as: str | None = None
) -> dict[str, Any]:
    """Fill {{key}} placeholders in a .docx template (`item` = absolute OneDrive path or item
    id). With `save_as` (absolute path) the result is a new file and the template is left
    untouched; otherwise the template is overwritten. Reports missing and leftover keys."""
    doc, meta = _load(item)
    counts = {k: _replace_all(doc, f"{{{{{k}}}}}", str(v)) for k, v in values.items()}
    leftover = sorted({m for p in _all_paragraphs(doc) for m in _PLACEHOLDER.findall(p.text)})
    result = {
        "replacements": sum(counts.values()),
        "missing": [k for k, n in counts.items() if not n],
        "leftover_placeholders": leftover,
    }
    if save_as:
        created = files.create_bytes(save_as, files.to_bytes(doc), "fail")
        return {**created, **result, "uploaded": True}
    changed = result["replacements"] > 0
    return files.save_if_changed(meta, changed, lambda: files.to_bytes(doc), **result)


@tool("files")
def word_export_pdf(item: str, dest: str | None = None, overwrite: bool = False) -> dict[str, Any]:
    """Convert a .docx (absolute OneDrive path or item id) to PDF and save it locally."""
    return files.files_convert(item, "pdf", dest, overwrite)
