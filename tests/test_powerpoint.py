import io

from pptx import Presentation
from pptx.util import Inches

from broker365.services import files, powerpoint
from tests.conftest import G

ITEM = {
    "id": "P!1",
    "name": "deck.pptx",
    "size": 100,
    "eTag": '"etag-p"',
    "file": {},
    "parentReference": {"path": "/drive/root:/Docs"},
}


def _deck(*titles: str):
    prs = Presentation()
    for t in titles:
        s = prs.slides.add_slide(prs.slide_layouts[1])
        s.shapes.title.text = t
        s.placeholders[1].text_frame.text = f"{t} body"
    return prs


def _serve(httpx_mock, prs, put: bool = True):
    httpx_mock.add_response(url=f"{G}/me/drive/root:/Docs/deck.pptx:", json=ITEM)
    httpx_mock.add_response(url=f"{G}/me/drive/items/P!1/content", content=files.to_bytes(prs))
    if put:
        httpx_mock.add_response(method="PUT", url=f"{G}/me/drive/items/P!1/content", json=ITEM)


def _uploaded(httpx_mock):
    put = [r for r in httpx_mock.get_requests() if r.method == "PUT"][-1]
    return Presentation(io.BytesIO(put.content)), put


def test_create_and_read_back(httpx_mock):
    httpx_mock.add_response(
        method="PUT",
        url=f"{G}/me/drive/root:/Docs/deck.pptx:/content?%40microsoft.graph.conflictBehavior=fail",
        json=ITEM,
    )
    out = powerpoint.ppt_create(
        "/Docs/deck.pptx",
        [
            {"title": "Kickoff", "subtitle": "Q3 plan"},
            {"title": "Goals", "bullets": ["Grow", "  - Revenue", "Hire"], "notes": "Say hi"},
            {"title": "End", "layout": "blank"},
        ],
    )
    assert out["slides"] == 3
    prs, _ = _uploaded(httpx_mock)
    s1, s2, s3 = prs.slides
    assert s1.slide_layout.name == "Title Slide"
    assert s1.placeholders[1].text_frame.text == "Q3 plan"
    paras = s2.placeholders[1].text_frame.paragraphs
    assert [(p.text, p.level) for p in paras] == [("Grow", 0), ("Revenue", 1), ("Hire", 0)]
    assert s2.notes_slide.notes_text_frame.text == "Say hi"
    assert s3.slide_layout.name == "Blank"
    assert any(sh.has_text_frame and sh.text_frame.text == "End" for sh in s3.shapes)

    _serve(httpx_mock, prs, put=False)
    read = powerpoint.ppt_read("/Docs/deck.pptx")
    assert read["name"] == "deck.pptx"
    assert read["slides"][1] == {
        "index": 2,
        "title": "Goals",
        "text": "Grow\nRevenue\nHire",
        "notes": "Say hi",
    }
    assert read["slides"][0]["text"] == "Q3 plan"


def test_add_slide_at_position(httpx_mock):
    _serve(httpx_mock, _deck("A", "B"))
    out = powerpoint.ppt_add_slide("/Docs/deck.pptx", "New", ["x"], notes="n", position=1)
    assert out["slide_index"] == 1 and out["slides"] == 3
    prs, put = _uploaded(httpx_mock)
    assert put.headers["If-Match"] == '"etag-p"'
    assert [s.shapes.title.text for s in prs.slides] == ["New", "A", "B"]
    assert prs.slides[0].notes_slide.notes_text_frame.text == "n"


def test_replace_text_shapes_and_tables(httpx_mock):
    prs = _deck("ACME intro")
    slide = prs.slides[0]
    tbl = slide.shapes.add_table(1, 1, Inches(1), Inches(5), Inches(3), Inches(1)).table
    tbl.cell(0, 0).text = "ACME"
    p = slide.placeholders[1].text_frame.add_paragraph()
    p.add_run().text = "AC"
    p.add_run().text = "ME!"
    _serve(httpx_mock, prs)
    out = powerpoint.ppt_replace_text("/Docs/deck.pptx", "ACME", "Globex")
    assert out["replacements"] == 4 and out["uploaded"]  # title, body, table, split run
    new, _ = _uploaded(httpx_mock)
    s = new.slides[0]
    assert s.shapes.title.text == "Globex intro"
    assert s.placeholders[1].text_frame.paragraphs[-1].text == "Globex!"
    table = next(sh for sh in s.shapes if sh.has_table).table
    assert table.cell(0, 0).text == "Globex"


def test_replace_zero_does_not_upload(httpx_mock):
    _serve(httpx_mock, _deck("A"), put=False)
    out = powerpoint.ppt_replace_text("/Docs/deck.pptx", "zzz", "y")
    assert out["replacements"] == 0 and out["uploaded"] is False
    assert all(r.method == "GET" for r in httpx_mock.get_requests())
