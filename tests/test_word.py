import io

import docx
import pytest

from broker365.errors import ConflictError
from broker365.services import files, word
from tests.conftest import G

ITEM = {
    "id": "W!1",
    "name": "report.docx",
    "size": 100,
    "eTag": '"etag-w"',
    "file": {},
    "parentReference": {"path": "/drive/root:/Docs"},
}


def _serve(httpx_mock, doc, put: bool = True, status: int = 200):
    httpx_mock.add_response(url=f"{G}/me/drive/root:/Docs/report.docx:", json=ITEM)
    httpx_mock.add_response(url=f"{G}/me/drive/items/W!1/content", content=files.to_bytes(doc))
    if put:
        httpx_mock.add_response(
            method="PUT",
            url=f"{G}/me/drive/items/W!1/content",
            status_code=status,
            json=ITEM
            if status == 200
            else {"error": {"code": "resourceModified", "message": "changed"}},
        )


def _uploaded(httpx_mock):
    put = [r for r in httpx_mock.get_requests() if r.method == "PUT"][-1]
    return docx.Document(io.BytesIO(put.content)), put


def test_create_from_markdown(httpx_mock):
    httpx_mock.add_response(
        method="PUT",
        url=f"{G}/me/drive/root:/Docs/new.docx:/content?%40microsoft.graph.conflictBehavior=fail",
        json={**ITEM, "name": "new.docx"},
    )
    md = """# Title

Intro with **bold** and *italic* and `code`.
continued line

- one
  - nested
1. first

| Name | Qty |
|------|-----|
| a    | 1   |

```
x = 1
```
"""
    out = word.word_create("/Docs/new.docx", md)
    assert out["name"] == "new.docx"
    doc, _ = _uploaded(httpx_mock)
    styles = [(p.style.name, p.text) for p in doc.paragraphs]
    assert ("Heading 1", "Title") in styles
    assert ("Normal", "Intro with bold and italic and code. continued line") in styles
    assert ("List Bullet", "one") in styles
    assert ("List Bullet 2", "nested") in styles
    assert ("List Number", "first") in styles
    intro = next(p for p in doc.paragraphs if p.text.startswith("Intro"))
    assert [r.text for r in intro.runs if r.bold] == ["bold"]
    assert [r.text for r in intro.runs if r.italic] == ["italic"]
    assert [r.text for r in intro.runs if r.font.name == word.MONO] == ["code"]
    code = next(p for p in doc.paragraphs if p.text == "x = 1")
    assert code.runs[0].font.name == word.MONO
    (table,) = doc.tables
    assert table.style.name == "Table Grid"
    assert [[c.text for c in r.cells] for r in table.rows] == [["Name", "Qty"], ["a", "1"]]
    assert table.cell(0, 0).paragraphs[0].runs[0].bold


def test_read(httpx_mock):
    doc = docx.Document()
    doc.add_heading("Hello", 1)
    doc.add_paragraph("Body")
    doc.add_paragraph("")
    t = doc.add_table(rows=1, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text = "a", "b"
    _serve(httpx_mock, doc, put=False)
    out = word.word_read("/Docs/report.docx")
    assert out["paragraphs"] == [
        {"style": "Heading 1", "text": "Hello"},
        {"style": "Normal", "text": "Body"},
    ]
    assert out["tables"] == [[["a", "b"]]]
    assert out["truncated"] is False
    assert out["path"] == "/Docs/report.docx"
    _serve(httpx_mock, doc, put=False)
    short = word.word_read("/Docs/report.docx", max_chars=6)
    assert short["truncated"] and short["tables"] == []
    assert [p["text"] for p in short["paragraphs"]] == ["Hello"]


def test_replace_across_runs_tables_and_headers(httpx_mock):
    doc = docx.Document()
    p = doc.add_paragraph()
    p.add_run("Dear ACME")
    p.add_run(" Corp").bold = True
    keep = doc.add_paragraph()
    keep.add_run("ACME Corp ")
    keep.add_run("rocks").italic = True
    doc.add_table(rows=1, cols=1).cell(0, 0).text = "Client: ACME Corp"
    doc.sections[0].header.paragraphs[0].text = "acme corp confidential"
    _serve(httpx_mock, doc)
    out = word.word_replace_text("/Docs/report.docx", "ACME Corp", "Globex", match_case=False)
    assert out["replacements"] == 4 and out["uploaded"]
    new, put = _uploaded(httpx_mock)
    assert put.headers["If-Match"] == '"etag-w"'
    assert new.paragraphs[0].text == "Dear Globex"
    assert [r.text for r in new.paragraphs[1].runs] == ["Globex ", "rocks"]
    assert new.paragraphs[1].runs[1].italic
    assert new.tables[0].cell(0, 0).text == "Client: Globex"
    assert new.sections[0].header.paragraphs[0].text == "Globex confidential"


def test_replace_match_case(httpx_mock):
    doc = docx.Document()
    doc.add_paragraph("acme and ACME")
    _serve(httpx_mock, doc)
    assert word.word_replace_text("/Docs/report.docx", "ACME", "X")["replacements"] == 1
    assert _uploaded(httpx_mock)[0].paragraphs[0].text == "acme and X"


def test_zero_replacements_does_not_upload(httpx_mock):
    doc = docx.Document()
    doc.add_paragraph("nothing here")
    _serve(httpx_mock, doc, put=False)
    out = word.word_replace_text("/Docs/report.docx", "absent", "x")
    assert out == {**out, "replacements": 0, "uploaded": False}
    assert all(r.method == "GET" for r in httpx_mock.get_requests())


def test_replace_conflict_propagates(httpx_mock):
    doc = docx.Document()
    doc.add_paragraph("old")
    _serve(httpx_mock, doc, status=412)
    with pytest.raises(ConflictError):
        word.word_replace_text("/Docs/report.docx", "old", "new")


def test_fill_template_save_as(httpx_mock):
    doc = docx.Document()
    p = doc.add_paragraph()
    for part in ("Hi {{", "name", "}}, total {{amount}} due {{date}}"):
        p.add_run(part)
    _serve(httpx_mock, doc, put=False)
    httpx_mock.add_response(
        method="PUT",
        url=f"{G}/me/drive/root:/Out/filled.docx:/content?%40microsoft.graph.conflictBehavior=fail",
        json={**ITEM, "id": "W!2", "name": "filled.docx"},
    )
    out = word.word_fill_template(
        "/Docs/report.docx",
        {"name": "Ann", "amount": "$5", "missing": "?"},
        save_as="/Out/filled.docx",
    )
    assert out["name"] == "filled.docx"
    assert out["replacements"] == 2
    assert out["missing"] == ["missing"]
    assert out["leftover_placeholders"] == ["date"]
    new, _ = _uploaded(httpx_mock)
    assert new.paragraphs[0].text == "Hi Ann, total $5 due {{date}}"


def test_append(httpx_mock):
    doc = docx.Document()
    doc.add_paragraph("existing")
    _serve(httpx_mock, doc)
    word.word_append("/Docs/report.docx", "## More\n\n- item")
    new, _ = _uploaded(httpx_mock)
    assert [(p.style.name, p.text) for p in new.paragraphs] == [
        ("Normal", "existing"),
        ("Heading 2", "More"),
        ("List Bullet", "item"),
    ]
